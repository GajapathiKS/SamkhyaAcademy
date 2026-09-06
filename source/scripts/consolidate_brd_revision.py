from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT.parent / "03_Specification_Documents"
SOURCE = SPEC / "SamkhyaAcademy_Functional_BRD.docx"
OUTPUT = SPEC / "SamkhyaAcademy_Functional_BRD_Consolidated_2026-09-02.docx"


def replace_paragraph_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    paragraph.text = paragraph.text.replace(old, new)
    return True


def add_bullet(document, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = document.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# Correct the legacy permission statements without rewriting unrelated role content.
roles = doc.tables[1]
for row in roles.rows[1:]:
    role = row.cells[0].text.strip()
    if role == "Platform Admin":
        row.cells[0].text = "Super Admin / Platform Admin"
        row.cells[2].text = (
            "Approve protected setting changes; publish immutable course versions; "
            "administer users, organizations, configuration, CRM oversight and content permissions."
        )
        row.cells[3].text = "Internal-only; requester cannot approve the same request."
    elif role == "Content Admin":
        row.cells[2].text = (
            "Create/edit/review courses, modules, lessons, exams, blogs, webinars, faculty, "
            "media and learning paths; submit drafts and protected-setting change requests."
        )
        row.cells[3].text = (
            "Cannot publish, approve own requests, or directly change protected core settings."
        )
    elif role == "Internal SME":
        row.cells[3].text = "Publishing and protected-setting approval require a Super Admin."

# Retire the standalone SAP offering while preserving SAP as an FDE integration example.
portfolio = doc.tables[4]
for row in portfolio.rows[1:]:
    if row.cells[0].text.strip() == "SAP Enterprise Consulting":
        values = [
            "Data Structures & Algorithms",
            "Pricing",
            "Full Online / Mixed Interactive",
            "Intermediate paid programming path",
            "Arrays, linked structures, trees, heaps, graphs, dynamic programming, SVG/HTML simulations and capstone.",
        ]
        for cell, value in zip(row.cells, values):
            cell.text = value

for paragraph in doc.paragraphs:
    replace_paragraph_text(
        paragraph,
        "19. Clickable prototype / Marvel mapping inventory (53 screens)",
        "19. Expanded clickable prototype / Marvel mapping inventory (95 screens)",
    )
    replace_paragraph_text(
        paragraph,
        "All five Venture Builder UX workflows are functional and persisted.",
        "The complete Venture Builder journey, continuation states, workspace, artifacts and mentor gates are functional and persisted.",
    )
    replace_paragraph_text(
        paragraph,
        "creates Marvel mapping manifest for the 53-screen prototype.",
        "maintains the generated 95-screen inventory and Marvel mapping manifests as the implementation-level acceptance index.",
    )

# Make the original 53-row table explicitly historical rather than silently misleading.
inventory = doc.tables[11]
inventory.cell(0, 1).text = "Original screen (historical baseline; superseded by 95-screen inventory)"
for row in inventory.rows[1:]:
    if row.cells[0].text.strip() == "13":
        row.cells[1].text = "RETIRED - SAP Enterprise Consulting"
        row.cells[2].text = "Archived source only"
        row.cells[3].text = "Excluded from implementation and acceptance"

doc.add_page_break()
heading = doc.add_paragraph("23. Consolidated 2026 learning-platform and UX requirements", style="Heading 1")
heading.runs[0].font.color.rgb = RGBColor(15, 35, 86)

lead = doc.add_paragraph()
lead_run = lead.add_run("Revision status: effective 2 September 2026. ")
lead_run.bold = True
lead.add_run(
    "This section consolidates the approved decisions developed during UX prototyping. "
    "It supersedes conflicting legacy screen counts, standalone SAP marketing, publishing permissions, "
    "and course-rendering assumptions while preserving all compatible requirements above."
)

doc.add_paragraph("23.1 Visual acceptance and golden masters", style="Heading 2")
add_bullet(doc, "The supplied approved screenshots and the 02_Approved_UX_References folder are visual acceptance authorities, not optional inspiration.")
add_bullet(doc, "Match composition, density, typography, Poppins-based hierarchy, indigo/navy/blue-violet palette, iconography, card geometry, spacing, borders, status colors and real-estate efficiency.")
add_bullet(doc, "Use real responsive HTML/CSS/SVG components. Never use a complete reference screenshot as a page background.")
add_bullet(doc, "The seven golden-master experiences are: Venture landing, Idea Intake, Venture Workflow, Artifacts & Launch Readiness, Mentor Review, Venture Workspace Overview and Unified Course Lesson.")
add_bullet(doc, "Content may change to remain BRD-accurate, but visual structure and interaction language must remain faithful to the approved masters.")
add_bullet(doc, "Each state must be captured at 1440px and checked for clipping, overflow, legibility, wasted real estate, responsive reflow and visual continuity.")
add_bullet(doc, "Empty real estate is a hard rejection condition: pages may not contain oversized blank bands, sparse desktop canvases, artificially tall cards or unused decision regions.")
add_bullet(doc, "Application states are captured at the customer viewport represented by the master; full-page capture is reserved for pages with continuous meaningful below-the-fold content.")
add_bullet(doc, "Dashboards use a consistent accessible line-icon system for navigation, KPI cards, search, filter, sort, columns, export, row actions, section headings, approvals and publishing; Unicode pseudo-icons are prohibited.")

doc.add_paragraph("23.2 Unified course and lesson experience", style="Heading 2")
add_bullet(doc, "One persistent course shell must support video, long-form text, mixed explanation/example, interactive code lab, SVG/HTML simulation, practice exercise, project package, quiz/exam and downloadable resources.")
add_bullet(doc, "The shell retains module navigation, progress, resume/next lesson, lesson overview, notes, discussion, assignments, resources and course facts across all lesson formats.")
add_bullet(doc, "Text-oriented learning must feel as complete as video learning: structured headings, callouts, runnable snippets, copy/export actions, practice checks and accessible code examples.")
add_bullet(doc, "Browser coding labs provide an editor, preview/output pane, reset/run controls, tests, hints, state persistence and mobile-safe fallback behavior.")
add_bullet(doc, "Project packages may include README/brief, assets, starter code, folder architecture, GIF/media, downloadable ZIP and completion evidence.")

doc.add_paragraph("23.3 Editorial blog boundary", style="Heading 2")
add_bullet(doc, "Blogs are public/editorial insight and discussion content, separate from courses and course completion.")
add_bullet(doc, "A blog may link to courses, webinars and calls to action, but must not create lesson progress, assessment completion or certificate eligibility.")
add_bullet(doc, "Course discussion remains attached to the relevant course or lesson and is not implemented as a blog post.")

doc.add_paragraph("23.4 Programming pathway and interactive simulations", style="Heading 2")
add_bullet(doc, "C Programming Fundamentals and C++ Essentials are free acquisition courses with tracked enrollment, text/video/mixed lessons, runnable examples, practice and downloadable starter projects.")
add_bullet(doc, "Data Structures & Algorithms is a paid intermediate course covering arrays/strings, linked structures, trees/heaps, graphs, dynamic programming and a problem-solving capstone.")
add_bullet(doc, "DSA lessons support code-native SVG/HTML node, edge and algorithm simulations with step controls, synchronized code, live state, explanations, checkpoints and accessibility labels.")
add_bullet(doc, "Simulation authoring stores semantic graph/state/step data and renderer configuration; it does not depend on flattened screenshots or videos.")

doc.add_paragraph("23.5 Content administration and publishing governance", style="Heading 2")
add_bullet(doc, "Multiple employees may hold the same role; a single employee may hold multiple approved roles. Shared employee logins are prohibited.")
add_bullet(doc, "Content Admin can create and edit all course-side content and submit drafts, reviews and protected-setting change requests.")
add_bullet(doc, "Content Admin cannot publish, approve their own request, or directly change protected fields such as pricing category, access type, delivery/content type, price, currency and preview entitlement.")
add_bullet(doc, "Super Admin approves or rejects protected-setting requests and publishes immutable course versions after quality and projection checks.")
add_bullet(doc, "Published learners remain pinned to the assigned immutable version unless an explicit migration is approved.")
add_bullet(doc, "Audit records identify the exact employee, action, before/after values, request, reviewer, decision, timestamp and published version.")

doc.add_paragraph("23.6 Expanded prototype and acceptance inventory", style="Heading 2")
add_bullet(doc, "The standalone SAP program screen is retired. SAP remains only as a factual enterprise-integration example inside FDE alongside SharePoint, CRM, Jira/ServiceNow-style systems and legacy platforms.")
add_bullet(doc, "The original 53-screen baseline becomes a renumbered 52-screen core after SAP retirement and is extended to a 95-screen customer prototype inventory.")
add_bullet(doc, "The expanded inventory includes multi-format learning, code labs, practice continuation states, blogs, project packages, Venture Builder continuation/workspace states, free C/C++ journeys, paid DSA simulations and two-person administration.")
add_bullet(doc, "docs/ux/SCREEN_INVENTORY.md and the generated Markdown/CSV Marvel maps are the implementation-level source for exact screen IDs, routes, contexts, statuses and screenshots.")
add_bullet(doc, "The inventory is extensible: additional continuation, empty, error, approval and responsive states may be added when required to demonstrate the complete customer experience.")

doc.add_paragraph("23.7 Acceptance additions", style="Heading 2")
add_bullet(doc, "Verify that video, text and mixed lesson types share entitlements, progress, comments, resources, exams and certificate behavior.")
add_bullet(doc, "Verify C and C++ free enrollment without payment and DSA locked/full states through paid or organization entitlement.")
add_bullet(doc, "Verify simulation step transitions, reset, checkpoints, keyboard access and saved progress without exposing protected answers.")
add_bullet(doc, "Verify Content Admin draft submission, protected-setting request, Super Admin approval/rejection, immutable publish and self-approval prevention.")
add_bullet(doc, "Verify that blogs never modify course progress and that standalone SAP marketing is absent from public discovery, brochures and prototype acceptance.")

doc.save(OUTPUT)
print(OUTPUT)
